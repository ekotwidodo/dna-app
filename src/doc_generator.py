"""Generator dokumen Laporan Harian dalam format PDF (ReportLab) dan Word (python-docx)."""

from __future__ import annotations

import io
import os
import re
import sys
from dataclasses import dataclass, field
from html import escape
from html.parser import HTMLParser
from pathlib import Path

from helpers import fit_size, format_tanggal, get_logo_path, get_signature_path

# --------------------------------------------------------------------------- #
# Model data laporan
# --------------------------------------------------------------------------- #


@dataclass
class ReportImage:
    data: bytes  # JPEG yang sudah dinormalisasi
    width: int
    height: int
    caption: str = ""


@dataclass
class ReportData:
    nama: str
    nip: str
    jabatan: str
    kegiatan: str  # Rencana kinerja / IKI / Nama kegiatan
    periode: str
    uraian: str  # Deskripsi kegiatan (mendukung RTF/HTML)
    hasil: str  # Deskripsi capaian (mendukung RTF/HTML)
    progres: int
    tanggal_ttd: object  # datetime.date
    kota: str = "Bandar Lampung"
    instansi: str = "BADAN PUSAT STATISTIK"
    wilayah: str = "PROVINSI LAMPUNG"
    link_dukung: str = ""
    use_logo: bool = True
    images: list[ReportImage] = field(default_factory=list)

    @property
    def hasil_lengkap(self) -> str:
        text = self.hasil.strip()
        prog_text = f"Progres: {self.progres}%"
        if "<" in text and ">" in text:
            return f"{text}<p><b>{prog_text}</b></p>"
        return f"{text}\n\n{prog_text}" if text else prog_text

    def rows(self) -> list[tuple[str, str]]:
        return [
            ("Nama", self.nama),
            ("NIP", self.nip),
            ("Jabatan", self.jabatan or "-"),
            ("Kegiatan", self.kegiatan),
            ("Periode", self.periode),
            ("Uraian Kegiatan", self.uraian),
            ("Hasil", self.hasil_lengkap),
        ]


# Ukuran (cm) yang dipakai bersama oleh PDF & DOCX
PAGE_MARGIN_CM = 2.0
CONTENT_WIDTH_CM = 21.0 - 2 * PAGE_MARGIN_CM  # A4 = 21 cm
LABEL_COL_CM = 4.2
LOGO_HEIGHT_CM = 2.4
IMG_BOX_W_CM = CONTENT_WIDTH_CM / 2 - 0.6
IMG_BOX_H_CM = IMG_BOX_W_CM * 0.75  # rasio 4:3


# --------------------------------------------------------------------------- #
# Konverter HTML / RTF sederhana
# --------------------------------------------------------------------------- #
def clean_html_for_reportlab(html_text: str) -> str:
    """Konversi HTML/RTF sederhana menjadi format yang didukung ReportLab Paragraph."""
    if not html_text:
        return "-"
    text = html_text.strip()
    if "<" not in text and ">" not in text:
        return escape(text).replace("\n", "<br/>")

    # Ubah strong/em menjadi b/i
    text = re.sub(r"</?strong[^>]*>", lambda m: "<b>" if m.group(0).startswith("<s") else "</b>", text, flags=re.I)
    text = re.sub(r"</?em[^>]*>", lambda m: "<i>" if m.group(0).startswith("<e") else "</i>", text, flags=re.I)

    # Format list unordered (ul -> bullet)
    def _repl_ul(match):
        items = re.findall(r"<li[^>]*>(.*?)</li>", match.group(1), flags=re.S | re.I)
        return "<br/>" + "".join(f"&bull; {it.strip()}<br/>" for it in items if it.strip())

    text = re.sub(r"<ul[^>]*>(.*?)</ul>", _repl_ul, text, flags=re.S | re.I)

    # Format list ordered (ol -> 1., 2., ...)
    def _repl_ol(match):
        items = re.findall(r"<li[^>]*>(.*?)</li>", match.group(1), flags=re.S | re.I)
        return "<br/>" + "".join(f"{idx+1}. {it.strip()}<br/>" for idx, it in enumerate(items) if it.strip())

    text = re.sub(r"<ol[^>]*>(.*?)</ol>", _repl_ol, text, flags=re.S | re.I)

    # Format block paragraph dan break
    text = re.sub(r"<br\s*/?>", "<br/>", text, flags=re.I)
    text = re.sub(r"</div>\s*<div[^>]*>", "<br/>", text, flags=re.I)
    text = re.sub(r"</p>\s*<p[^>]*>", "<br/><br/>", text, flags=re.I)
    text = re.sub(r"</p>", "<br/>", text, flags=re.I)
    text = re.sub(r"</div>", "<br/>", text, flags=re.I)
    text = re.sub(r"</?(?:p|div|span)[^>]*>", "", text, flags=re.I)

    # Buang tag lain selain b, i, u, br
    text = re.sub(r"<(?!(?:/?(?:b|i|u|br\s*/?)))[^>]+>", "", text, flags=re.I)
    text = re.sub(r"(?:<br\s*/?>\s*){3,}", "<br/><br/>", text, flags=re.I)
    text = re.sub(r"^(?:<br\s*/?>\s*)+|(?:<br\s*/?>\s*)+$", "", text, flags=re.I)
    return text.strip() or "-"



class DocxHTMLParser(HTMLParser):
    """Parser HTML sederhana untuk membuat paragraf dan run berformat di python-docx."""

    def __init__(self, cell):
        super().__init__()
        self.cell = cell
        self.current_p = cell.paragraphs[0]
        self.current_p.paragraph_format.line_spacing = 1.15
        self.is_first_p = True
        self.bold = False
        self.italic = False
        self.underline = False
        self.list_type = None
        self.list_index = 0

    def new_paragraph(self):
        if self.is_first_p and len(self.current_p.text) == 0:
            self.is_first_p = False
            self.current_p.paragraph_format.line_spacing = 1.0
            return self.current_p
        self.is_first_p = False
        self.current_p = self.cell.add_paragraph()
        self.current_p.paragraph_format.line_spacing = 1.0
        return self.current_p

    def handle_starttag(self, tag, attrs):
        t = tag.lower()
        if t in ("b", "strong"):
            self.bold = True
        elif t in ("i", "em"):
            self.italic = True
        elif t == "u":
            self.underline = True
        elif t in ("p", "div"):
            if not self.is_first_p or len(self.current_p.text) > 0:
                self.new_paragraph()
        elif t == "ul":
            self.list_type = "ul"
        elif t == "ol":
            self.list_type = "ol"
            self.list_index = 0
        elif t == "li":
            self.new_paragraph()
            prefix = f"{self.list_index + 1}. " if self.list_type == "ol" else "• "
            if self.list_type == "ol":
                self.list_index += 1
            run = self.current_p.add_run(prefix)
            run.bold = True
        elif t == "br":
            self.current_p.add_break()

    def handle_endtag(self, tag):
        t = tag.lower()
        if t in ("b", "strong"):
            self.bold = False
        elif t in ("i", "em"):
            self.italic = False
        elif t == "u":
            self.underline = False
        elif t in ("ul", "ol"):
            self.list_type = None

    def handle_data(self, data):
        if not data:
            return
        run = self.current_p.add_run(data)
        run.bold = self.bold
        run.italic = self.italic
        run.underline = self.underline


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _find_font(candidates: list[str]) -> Path | None:
    """Mencari file font di berbagai direktori standar (Windows, macOS, Linux)."""
    search_dirs: list[Path] = []
    if os.name == "nt":
        search_dirs.append(Path("C:/Windows/Fonts"))
        windir = os.environ.get("WINDIR")
        if windir:
            search_dirs.append(Path(windir) / "Fonts")
    elif sys.platform == "darwin":
        search_dirs.extend([
            Path("/System/Library/Fonts/Supplemental"),
            Path("/System/Library/Fonts"),
            Path("/Library/Fonts"),
            Path.home() / "Library/Fonts",
        ])
    else:  # Linux & BSD
        search_dirs.extend([
            Path("/usr/share/fonts"),
            Path("/usr/local/share/fonts"),
            Path.home() / ".fonts",
            Path.home() / ".local/share/fonts",
        ])

    for d in search_dirs:
        if not d.exists():
            continue
        for name in candidates:
            exact = d / name
            if exact.is_file():
                return exact
            # Pencarian rekursif jika belum ketemu
            try:
                for match in d.rglob(name):
                    if match.is_file():
                        return match
            except Exception:
                pass
    return None


def _register_fonts() -> dict[str, str]:
    """Daftarkan Arial untuk kop dan Calibri untuk seluruh isi laporan (kompatibel lintas OS: Windows, macOS, Linux)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # Arial (khusus kop)
    reg_a = _find_font(["arial.ttf", "Arial.ttf", "LiberationSans-Regular.ttf", "DejaVuSans.ttf"])
    bd_a = _find_font(["arialbd.ttf", "Arial Bold.ttf", "Arial-Bold.ttf", "LiberationSans-Bold.ttf", "DejaVuSans-Bold.ttf"])
    bi_a = _find_font(["arialbi.ttf", "Arial Bold Italic.ttf", "Arial-BoldItalic.ttf", "LiberationSans-BoldItalic.ttf", "DejaVuSans-BoldOblique.ttf"])

    if reg_a and bd_a:
        try:
            if "Arial" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Arial", str(reg_a)))
                pdfmetrics.registerFont(TTFont("Arial-Bold", str(bd_a)))
            if bi_a and "Arial-BoldItalic" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Arial-BoldItalic", str(bi_a)))
            kop_font = "Arial-BoldItalic" if (bi_a and "Arial-BoldItalic" in pdfmetrics.getRegisteredFontNames()) else "Arial-Bold"
        except Exception:
            kop_font = "Helvetica-BoldOblique"
    else:
        kop_font = "Helvetica-BoldOblique"

    # Calibri (untuk isi dokumen: judul, label, isi, tabel, tanda tangan)
    cal_reg = _find_font(["calibri.ttf", "Calibri.ttf", "Carlito-Regular.ttf", "LiberationSans-Regular.ttf", "DejaVuSans.ttf"])
    cal_bd = _find_font(["calibrib.ttf", "Calibri Bold.ttf", "Calibri-Bold.ttf", "Carlito-Bold.ttf", "LiberationSans-Bold.ttf", "DejaVuSans-Bold.ttf"])
    cal_it = _find_font(["calibrii.ttf", "Calibri Italic.ttf", "Calibri-Italic.ttf", "Carlito-Italic.ttf", "LiberationSans-Italic.ttf", "DejaVuSans-Oblique.ttf"])
    cal_bi = _find_font(["calibriz.ttf", "Calibri Bold Italic.ttf", "Calibri-BoldItalic.ttf", "Carlito-BoldItalic.ttf", "LiberationSans-BoldItalic.ttf", "DejaVuSans-BoldOblique.ttf"])

    if cal_reg and cal_bd:
        try:
            if "Calibri" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Calibri", str(cal_reg)))
                pdfmetrics.registerFont(TTFont("Calibri-Bold", str(cal_bd)))
            if cal_it and "Calibri-Italic" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Calibri-Italic", str(cal_it)))
            if cal_bi and "Calibri-BoldItalic" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Calibri-BoldItalic", str(cal_bi)))
            body_font = "Calibri"
            body_bold = "Calibri-Bold"
        except Exception:
            body_font = "Helvetica"
            body_bold = "Helvetica-Bold"
    else:
        body_font = "Helvetica"
        body_bold = "Helvetica-Bold"

    return {
        "kop": kop_font,
        "body": body_font,
        "body_bold": body_bold,
    }


def build_pdf(data: ReportData, out_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable,
        Image,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    f = _register_fonts()
    st_kop = ParagraphStyle(
        "kop_compact",
        fontName=f["kop"],
        fontSize=11,
        leading=14.5,
        alignment=TA_LEFT,
    )
    st_title = ParagraphStyle("title", fontName=f["body_bold"], fontSize=14, leading=18, alignment=TA_CENTER)
    st_label = ParagraphStyle("label", fontName=f["body_bold"], fontSize=11, leading=15.2)
    st_body = ParagraphStyle("body", fontName=f["body"], fontSize=11, leading=15.2, alignment=TA_LEFT)
    st_caption = ParagraphStyle(
        "cap",
        fontName=f["body"],
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#374151"),
    )
    st_sign = ParagraphStyle("sign", fontName=f["body"], fontSize=11, leading=15.2, alignment=TA_CENTER)
    st_sign_name = ParagraphStyle("signname", parent=st_sign, fontName=f["body_bold"], fontSize=11, leading=15.2)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=PAGE_MARGIN_CM * cm,
        rightMargin=PAGE_MARGIN_CM * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=f"Laporan Harian - {data.nama}",
        author=data.nama,
        subject=data.periode,
    )
    story: list = []

    # ---- Kop (1 Row: Logo di kiri, Teks Bold Italic Uppercase di kanan) ----
    logo_path = get_logo_path()
    instansi_clean = data.instansi.strip().upper()
    wilayah_clean = data.wilayah.strip().upper()
    kop_text = f"<i><b>{escape(instansi_clean)}</b></i><br/><i><b>{escape(wilayah_clean)}</b></i>"

    use_logo = getattr(data, "use_logo", True)
    if logo_path.exists() and use_logo:
        from PIL import Image as PILImage

        with PILImage.open(logo_path) as logo:
            lw, lh = logo.size
        # Hitung ukuran logo proporsional agar cocok untuk rasio apapun (persegi, memanjang, emblem)
        w_raw, h_raw = fit_size(lw, lh, box_w=2.4 * cm, box_h=1.35 * cm)
        logo_img = Image(str(logo_path), width=w_raw, height=h_raw, hAlign="LEFT")
        total_w = CONTENT_WIDTH_CM * cm
        col_gap = 0.35 * cm
        col_logo_w = w_raw + col_gap
        kop_table = Table(
            [[logo_img, Paragraph(kop_text, st_kop)]],
            colWidths=[col_logo_w, total_w - col_logo_w],
        )
        kop_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))
        story.append(kop_table)
    else:
        story.append(Paragraph(kop_text, st_kop))

    # Di bawah kop: tanpa garis, langsung spasi dan judul LAPORAN HARIAN 14pt
    story.append(Spacer(1, 0.75 * cm))
    story.append(Paragraph("LAPORAN HARIAN", st_title))
    story.append(Spacer(1, 0.3 * cm))

    # ---- Tabel isi (2 kolom, tanpa titik dua) ----
    grid = colors.HexColor("#9AA3AF")
    shade = colors.HexColor("#F1F4F8")
    value_w = CONTENT_WIDTH_CM - LABEL_COL_CM

    rows = []
    for label, value in data.rows():
        formatted_val = clean_html_for_reportlab(value)
        rows.append([
            Paragraph(escape(label), st_label),
            Paragraph(formatted_val, st_body),
        ])
    dok_val = ""
    if data.link_dukung:
        # Tampilan mirip tombol dengan background warna biru (membutuhkan support reportlab 3.x)
        dok_val = f'<a href="{escape(data.link_dukung)}" color="white"><font backColor="#2563eb"> &nbsp; <b>BUKA TAUTAN</b> &nbsp; </font></a>'
    elif not data.images:
        dok_val = "-"
    rows.append([Paragraph("Dokumentasi", st_label), Paragraph(dok_val, st_body) if dok_val else ""])

    info = Table(rows, colWidths=[LABEL_COL_CM * cm, value_w * cm])
    info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.6, grid),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, grid),
        ("LINEAFTER", (0, 0), (0, -1), 0.4, grid),  # Garis pemisah kolom label & nilai
        ("BACKGROUND", (0, 0), (0, -1), shade),     # Shading kolom label
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(info)

    # ---- Dokumentasi: grid 2 kolom dengan caption kustom ----
    if data.images:
        img_rows: list[list] = []
        for i in range(0, len(data.images), 2):
            row = []
            for j, img in enumerate(data.images[i:i + 2]):
                w, h = fit_size(img.width, img.height, IMG_BOX_W_CM * cm, IMG_BOX_H_CM * cm)
                caption_text = img.caption.strip() if img.caption and img.caption.strip() else f"Gambar {i + j + 1}"
                cell = Table(
                    [
                        [Image(io.BytesIO(img.data), width=w, height=h)],
                        [Paragraph(escape(caption_text), st_caption)],
                    ],
                    rowHeights=[IMG_BOX_H_CM * cm + 4, None],
                )
                cell.setStyle(TableStyle([
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                    ("TOPPADDING", (0, 1), (-1, 1), 3),
                    ("BOTTOMPADDING", (0, 1), (-1, 1), 4),
                ]))
                row.append(cell)
            if len(row) == 1:
                row.append("")
            img_rows.append(row)

        gallery = Table(img_rows, colWidths=[CONTENT_WIDTH_CM / 2 * cm] * 2)
        gallery.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.6, grid),
            ("LINEAFTER", (0, 0), (0, -1), 0.4, grid),
            ("LINEBELOW", (0, 0), (-1, -2), 0.4, grid),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(gallery)

    # ---- Tanda tangan ----
    sig_path = get_signature_path()
    if sig_path.exists():
        from PIL import Image as PILImage

        with PILImage.open(sig_path) as sig:
            sw, sh = sig.size
        sig_w, sig_h = fit_size(sw, sh, 4.5 * cm, 2.0 * cm)
        sign_space = Image(str(sig_path), width=sig_w, height=sig_h)
    else:
        sign_space = Spacer(1, 2.0 * cm)

    sign_block = [
        [Paragraph(f"{escape(data.kota)}, {format_tanggal(data.tanggal_ttd)}", st_sign)],
        [Paragraph("Pegawai yang melaporkan,", st_sign)],
        [sign_space],
        [Paragraph(f"{escape(data.nama)}", st_sign_name)],
    ]
    sign_table = Table(sign_block, colWidths=[7.5 * cm])
    sign_table.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    wrapper = Table([["", sign_table]], colWidths=[(CONTENT_WIDTH_CM - 7.5) * cm, 7.5 * cm])
    wrapper.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(KeepTogether([Spacer(1, 0.8 * cm), wrapper]))

    doc.build(story)
    return out_path


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _set_cell_shading(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_table_borders(table, size: int = 4, color: str = "9AA3AF", none: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil" if none else "single")
        if not none:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _add_bottom_border(paragraph, size: int = 18) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "thinThickSmallGap" if size > 12 else "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    bdr.append(bottom)
    p_pr.append(bdr)


def _write_multiline(paragraph, text: str, bold: bool = False) -> None:
    paragraph.paragraph_format.line_spacing = 1.0
    lines = (text.strip() or "-").split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.bold = bold
        if i < len(lines) - 1:
            run.add_break()


def _add_docx_button_link(paragraph, url: str, text: str = "BUKA TAUTAN"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    
    # Text Color (White)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "FFFFFF")
    rPr.append(c)

    # Bold
    b = OxmlElement("w:b")
    rPr.append(b)

    # Background shading (Blue)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "2563EB")
    rPr.append(shd)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = f"   {text}   "
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _write_value_to_docx_cell(cell, value: str) -> None:
    text = value.strip()
    if "<" in text and ">" in text:
        parser = DocxHTMLParser(cell)
        parser.feed(text)
    else:
        _write_multiline(cell.paragraphs[0], text)


def build_docx(data: ReportData, out_path: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(PAGE_MARGIN_CM)
    section.top_margin = section.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    core = doc.core_properties
    core.title = f"Laporan Harian - {data.nama}"
    core.author = data.nama
    core.subject = data.periode

    def centered(text: str = "", size: float = 11, bold: bool = True, space_after: float = 0, font_name: str = "Calibri"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.name = font_name
            r.font.size = Pt(size)
        return p

    # ---- Kop (1 Row: Logo di kiri, Teks Bold Italic Uppercase di kanan) ----
    logo_path = get_logo_path()
    instansi_clean = data.instansi.strip().upper()
    wilayah_clean = data.wilayah.strip().upper()

    kop_table = doc.add_table(rows=1, cols=2)
    kop_table.autofit = False
    _set_table_borders(kop_table, none=True)
    c_left, c_right = kop_table.rows[0].cells
    c_left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c_right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    use_logo = getattr(data, "use_logo", True)
    if logo_path.exists() and use_logo:
        from PIL import Image as PILImage
        with PILImage.open(logo_path) as limg:
            lw, lh = limg.size
        w_cm, h_cm = fit_size(lw, lh, box_w=2.2, box_h=1.35)
        logo_col_w = max(w_cm + 0.3, 1.8)
        c_left.width = Cm(logo_col_w)
        c_right.width = Cm(CONTENT_WIDTH_CM - logo_col_w)
        p_img = c_left.paragraphs[0]
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(0)
        p_img.add_run().add_picture(str(logo_path), width=Cm(w_cm), height=Cm(h_cm))
    else:
        c_left.width = Cm(0.01)
        c_right.width = Cm(CONTENT_WIDTH_CM)

    p_txt1 = c_right.paragraphs[0]
    p_txt1.paragraph_format.space_before = Pt(0)
    p_txt1.paragraph_format.space_after = Pt(1)
    r1 = p_txt1.add_run(instansi_clean)
    r1.bold = True
    r1.italic = True
    r1.font.name = "Arial"
    r1.font.size = Pt(11)

    p_txt2 = c_right.add_paragraph()
    p_txt2.paragraph_format.space_before = Pt(0)
    p_txt2.paragraph_format.space_after = Pt(0)
    r2 = p_txt2.add_run(wilayah_clean)
    r2.bold = True
    r2.italic = True
    r2.font.name = "Arial"
    r2.font.size = Pt(11)

    # Garis pemisah kop dihilangkan, beri sedikit spasi kosong
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    p_space.paragraph_format.space_after = Pt(12)

    # Judul (14 pt, bold, font Calibri)
    p_title = centered("LAPORAN HARIAN", size=14, bold=True, space_after=10, font_name="Calibri")

    # ---- Tabel isi (2 kolom, tanpa titik dua) ----
    value_w = CONTENT_WIDTH_CM - LABEL_COL_CM
    widths = [Cm(LABEL_COL_CM), Cm(value_w)]
    rows = data.rows() + [("Dokumentasi", "")]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)

    for r, (label, value) in enumerate(rows):
        cells = table.rows[r].cells
        cells[0].width = widths[0]
        cells[1].width = widths[1]
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        _write_multiline(cells[0].paragraphs[0], label, bold=True)
        _set_cell_shading(cells[0], "F8FAFC")

        if label != "Dokumentasi":
            _write_value_to_docx_cell(cells[1], value)
        else:
            if data.link_dukung:
                _add_docx_button_link(cells[1].paragraphs[0], data.link_dukung)
            elif not data.images:
                _write_multiline(cells[1].paragraphs[0], "-")

    # ---- Dokumentasi: grid 2 kolom dengan caption kustom ----
    if data.images:
        n_rows = (len(data.images) + 1) // 2
        gallery = doc.add_table(rows=n_rows, cols=2)
        gallery.alignment = WD_TABLE_ALIGNMENT.CENTER
        gallery.autofit = False
        _set_table_borders(gallery)
        for idx, img in enumerate(data.images):
            cell = gallery.rows[idx // 2].cells[idx % 2]
            cell.width = Cm(CONTENT_WIDTH_CM / 2)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            w, h = fit_size(img.width, img.height, IMG_BOX_W_CM, IMG_BOX_H_CM)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(io.BytesIO(img.data), width=Cm(w), height=Cm(h))

            caption_text = img.caption.strip() if img.caption and img.caption.strip() else f"Gambar {idx + 1}"
            cap = cell.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(6)
            run = cap.add_run(caption_text)
            run.font.size = Pt(8.5)

    # ---- Tanda tangan ----
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    sign = doc.add_table(rows=1, cols=2)
    sign.autofit = False
    _set_table_borders(sign, none=True)
    left, right = sign.rows[0].cells
    left.width = Cm(CONTENT_WIDTH_CM - 7.5)
    right.width = Cm(7.5)

    def sign_line(text: str, bold: bool = False, underline: bool = False, first: bool = False):
        p = right.paragraphs[0] if first else right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.underline = bold, underline
        return p

    sign_line(f"{data.kota}, {format_tanggal(data.tanggal_ttd)}", first=True)
    sign_line("Pegawai yang melaporkan,")
    sig_path = get_signature_path()
    if sig_path.exists():
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(sig_path), height=Cm(2.0))
    else:
        spacer = right.add_paragraph()
        spacer.paragraph_format.space_after = Pt(48)
    sign_line(data.nama, bold=True, underline=False)

    doc.save(str(out_path))
    return out_path
